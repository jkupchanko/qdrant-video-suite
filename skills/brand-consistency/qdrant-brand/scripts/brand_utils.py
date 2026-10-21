"""
Qdrant Brand Utilities
======================
Python utility module for applying Qdrant brand guidelines programmatically.
Import this module when creating branded presentations, documents, spreadsheets,
or HTML assets.

Usage:
    from brand_utils import QdrantBrand, QdrantPptxHelper, QdrantDocxHelper, QdrantXlsxHelper

Dependencies:
    - python-pptx (for presentations)
    - python-docx (for Word documents)
    - openpyxl (for spreadsheets)

All dependencies are optional. Only import the helper you need.
"""

# =============================================================================
# BRAND CONSTANTS
# =============================================================================

class QdrantColors:
    """Complete Qdrant color palette. All values are hex strings."""

    # Primary (Qdrant Red)
    PRIMARY_50 = "#FDECEF"
    PRIMARY_100 = "#FCD7DD"
    PRIMARY_200 = "#FAAEB3"
    PRIMARY_300 = "#FF8792"
    PRIMARY_400 = "#FF516B"
    PRIMARY_500 = "#DC244C"  # Main brand red
    PRIMARY_600 = "#CC1845"  # Hover
    PRIMARY_700 = "#BD0C3E"  # Active
    PRIMARY_800 = "#9E0D38"
    PRIMARY_900 = "#850D35"

    # Secondary (Violet)
    SECONDARY_50 = "#E5E8FF"
    SECONDARY_100 = "#DDE0FF"
    SECONDARY_200 = "#C2C5FF"
    SECONDARY_300 = "#9494FF"
    SECONDARY_400 = "#7E75FF"
    SECONDARY_500 = "#6047FF"  # Main secondary
    SECONDARY_600 = "#5F36F5"
    SECONDARY_700 = "#522AD8"
    SECONDARY_800 = "#4325AE"
    SECONDARY_900 = "#392689"

    # Neutrals
    WHITE = "#FFFFFF"
    NEUTRAL_50 = "#F0F3FA"   # Primary text on dark
    NEUTRAL_100 = "#E2E7F5"
    NEUTRAL_200 = "#D3D9EB"  # Light borders
    NEUTRAL_300 = "#CAD0E5"
    NEUTRAL_400 = "#ABB1C7"  # Disabled text
    NEUTRAL_500 = "#898EA3"
    NEUTRAL_600 = "#656B7F"  # Secondary text
    NEUTRAL_700 = "#4E5366"  # Dark borders
    NEUTRAL_800 = "#303547"  # Document body text
    NEUTRAL_900 = "#111824"  # Surface elevation 1
    NEUTRAL_950 = "#0B0F19"  # Dark background
    BLACK = "#000000"

    # Backgrounds (Dark Theme)
    BG_DEFAULT = "#0B0F19"
    BG_ELEVATION_1 = "#111824"
    BG_ELEVATION_2 = "#141A2A"
    BG_ELEVATION_8 = "#212635"
    BG_ELEVATION_24 = "#2B303E"

    # Backgrounds (Light Theme)
    BG_LIGHT_DEFAULT = "#FFFFFF"
    BG_LIGHT_SURFACE = "#F7F8FA"

    # Semantic
    SUCCESS = "#00897B"
    ERROR = "#D32F2F"
    WARNING = "#EF6C00"
    INFO = "#29B6F6"

    # Chart colors (ordered sequence)
    CHART_SEQUENCE = [
        "#DC244C",  # Red
        "#6047FF",  # Violet
        "#009688",  # Teal
        "#03A9F4",  # Light Blue
        "#FF9800",  # Orange
        "#9C27B0",  # Purple
        "#4CAF50",  # Green
        "#00BCD4",  # Cyan
    ]


class QdrantTypography:
    """Typography constants."""

    FONT_FAMILY = "Mona Sans"
    FONT_FAMILY_FALLBACK = '"Mona Sans", "Inter", -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif'
    FONT_MONO = "Geist Mono"
    FONT_MONO_FALLBACK = '"Geist Mono", "JetBrains Mono", "Fira Code", "SF Mono", Consolas, monospace'

    # Fallback font for Office documents (when Mona Sans isn't installed)
    FONT_OFFICE_FALLBACK = "Calibri"

    WEIGHT_REGULAR = 400
    WEIGHT_MEDIUM = 500
    WEIGHT_SEMIBOLD = 600
    WEIGHT_BOLD = 700

    # Type scale (size_px, weight, line_height_ratio, letter_spacing_px)
    SCALE = {
        "display-2": (104, 700, 1.05, -1.8),
        "display-1": (80, 600, 1.05, -1.8),
        "h1": (64, 600, 1.1, -1.5),
        "h2": (56, 600, 1.1, -0.5),
        "h3": (48, 600, 1.1, -0.5),
        "h4": (40, 600, 1.2, -0.5),
        "h5": (32, 600, 1.3, -0.5),
        "h6": (24, 600, 1.3, -0.5),
        "text-xl": (20, 400, 1.5, -0.2),
        "text-lg": (18, 400, 1.5, 0),
        "text-md": (16, 400, 1.5, 0),
        "text-sm": (14, 400, 1.5, 0),
        "text-xs": (12, 400, 1.5, 0),
    }

    # Presentation-specific sizes (pt)
    PPTX_TITLE = 44
    PPTX_SECTION = 36
    PPTX_HEADING = 28
    PPTX_SUBHEADING = 22
    PPTX_BODY = 18
    PPTX_CAPTION = 14
    PPTX_FOOTNOTE = 12

    # Document-specific sizes (pt)
    DOCX_TITLE = 28
    DOCX_H1 = 22
    DOCX_H2 = 18
    DOCX_H3 = 14
    DOCX_BODY = 11
    DOCX_CAPTION = 9


class QdrantSpacing:
    """Spacing and layout constants."""

    GRID_UNIT = 8  # Base unit in px

    # Common spacing values (px)
    XS = 8
    SM = 16
    MD = 24
    LG = 32
    XL = 40
    XXL = 48
    SECTION = 64
    SECTION_LG = 80
    HERO = 88

    # Layout
    MAX_WIDTH = 1110  # px
    SIDE_MARGIN = 165  # px (at 1440px viewport)
    COLUMN_GAP = 30  # px
    BORDER_RADIUS = 8  # px
    BORDER_RADIUS_LG = 16  # px

    # Slide layout (inches)
    SLIDE_WIDTH = 13.333
    SLIDE_HEIGHT = 7.5
    SLIDE_MARGIN_LEFT = 0.75
    SLIDE_MARGIN_RIGHT = 0.75
    SLIDE_MARGIN_TOP = 1.0
    SLIDE_MARGIN_BOTTOM = 0.5


class QdrantBrand:
    """Convenience class bundling all brand constants."""
    colors = QdrantColors
    typography = QdrantTypography
    spacing = QdrantSpacing


# =============================================================================
# UTILITY FUNCTIONS
# =============================================================================

def hex_to_rgb(hex_color: str) -> tuple:
    """Convert hex color string to RGB tuple."""
    hex_color = hex_color.lstrip('#')
    return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))


def hex_to_pptx_rgb(hex_color: str):
    """Convert hex color string to python-pptx RGBColor."""
    from pptx.dml.color import RGBColor
    r, g, b = hex_to_rgb(hex_color)
    return RGBColor(r, g, b)


def hex_to_docx_rgb(hex_color: str):
    """Convert hex color string to python-docx RGBColor."""
    from docx.shared import RGBColor
    r, g, b = hex_to_rgb(hex_color)
    return RGBColor(r, g, b)


def hex_to_openpyxl(hex_color: str) -> str:
    """Convert hex color to openpyxl-compatible string (no #)."""
    return hex_color.lstrip('#')


def get_chart_color(index: int) -> str:
    """Get chart color by index (wraps around if > 8)."""
    return QdrantColors.CHART_SEQUENCE[index % len(QdrantColors.CHART_SEQUENCE)]


def generate_css_variables(theme: str = "dark") -> str:
    """Generate CSS custom properties block for Qdrant theme.

    Args:
        theme: 'dark' or 'light'

    Returns:
        CSS string with custom properties
    """
    c = QdrantColors
    t = QdrantTypography
    s = QdrantSpacing

    if theme == "dark":
        bg_default = c.BG_DEFAULT
        bg_surface = c.BG_ELEVATION_1
        bg_surface_2 = c.BG_ELEVATION_2
        text_primary = c.NEUTRAL_50
        text_secondary = c.NEUTRAL_600
        border_color = c.NEUTRAL_700
    else:
        bg_default = c.BG_LIGHT_DEFAULT
        bg_surface = c.BG_LIGHT_SURFACE
        bg_surface_2 = c.NEUTRAL_50
        text_primary = c.NEUTRAL_950
        text_secondary = c.NEUTRAL_600
        border_color = c.NEUTRAL_200

    return f""":root {{
  --qdrant-primary: {c.PRIMARY_500};
  --qdrant-primary-hover: {c.PRIMARY_600};
  --qdrant-primary-active: {c.PRIMARY_700};
  --qdrant-secondary: {c.SECONDARY_500};
  --qdrant-bg-default: {bg_default};
  --qdrant-bg-surface: {bg_surface};
  --qdrant-bg-surface-2: {bg_surface_2};
  --qdrant-text-primary: {text_primary};
  --qdrant-text-secondary: {text_secondary};
  --qdrant-border: {border_color};
  --qdrant-success: {c.SUCCESS};
  --qdrant-error: {c.ERROR};
  --qdrant-warning: {c.WARNING};
  --qdrant-info: {c.INFO};
  --qdrant-font-family: {t.FONT_FAMILY_FALLBACK};
  --qdrant-font-mono: {t.FONT_MONO_FALLBACK};
  --qdrant-border-radius: {s.BORDER_RADIUS}px;
  --qdrant-max-width: {s.MAX_WIDTH}px;
  --qdrant-gradient: linear-gradient(135deg, {c.PRIMARY_500}, {c.SECONDARY_500});
}}"""


# =============================================================================
# PPTX HELPER
# =============================================================================

class QdrantPptxHelper:
    """Helper for creating Qdrant-branded presentations."""

    def __init__(self):
        from pptx import Presentation
        from pptx.util import Inches, Pt

        self.prs = Presentation()
        self.prs.slide_width = Inches(QdrantSpacing.SLIDE_WIDTH)
        self.prs.slide_height = Inches(QdrantSpacing.SLIDE_HEIGHT)

    def add_dark_slide(self):
        """Add a blank slide with dark background."""
        from pptx.util import Inches
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = hex_to_pptx_rgb(QdrantColors.BG_DEFAULT)
        return slide

    def add_title_slide(self, title: str, subtitle: str = ""):
        """Add a branded title slide with accent line."""
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN

        slide = self.add_dark_slide()

        # Accent line
        line = slide.shapes.add_shape(
            1, Inches(0.75), Inches(2.2), Inches(2), Pt(4)
        )
        line.fill.solid()
        line.fill.fore_color.rgb = hex_to_pptx_rgb(QdrantColors.PRIMARY_500)
        line.line.fill.background()

        # Title
        txBox = slide.shapes.add_textbox(
            Inches(0.75), Inches(2.5), Inches(11.5), Inches(2)
        )
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(QdrantTypography.PPTX_TITLE)
        p.font.color.rgb = hex_to_pptx_rgb(QdrantColors.NEUTRAL_50)
        p.font.name = QdrantTypography.FONT_OFFICE_FALLBACK
        p.alignment = PP_ALIGN.LEFT

        if subtitle:
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.font.size = Pt(QdrantTypography.PPTX_BODY)
            p2.font.color.rgb = hex_to_pptx_rgb(QdrantColors.NEUTRAL_600)
            p2.font.name = QdrantTypography.FONT_OFFICE_FALLBACK
            p2.space_before = Pt(24)

        return slide

    def add_content_slide(self, heading: str, body: str = ""):
        """Add a content slide with heading and body text."""
        from pptx.util import Inches, Pt

        slide = self.add_dark_slide()

        # Heading
        h = slide.shapes.add_textbox(
            Inches(0.75), Inches(0.75), Inches(11.5), Inches(0.8)
        )
        tf = h.text_frame
        p = tf.paragraphs[0]
        p.text = heading
        p.font.size = Pt(QdrantTypography.PPTX_HEADING)
        p.font.color.rgb = hex_to_pptx_rgb(QdrantColors.NEUTRAL_50)
        p.font.name = QdrantTypography.FONT_OFFICE_FALLBACK

        if body:
            b = slide.shapes.add_textbox(
                Inches(0.75), Inches(1.8), Inches(11.5), Inches(4.5)
            )
            tf = b.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = body
            p.font.size = Pt(QdrantTypography.PPTX_BODY)
            p.font.color.rgb = hex_to_pptx_rgb(QdrantColors.NEUTRAL_600)
            p.font.name = QdrantTypography.FONT_OFFICE_FALLBACK

        return slide

    def add_card(self, slide, left_inches, top_inches, width_inches, height_inches):
        """Add a branded card shape to a slide."""
        from pptx.util import Inches, Pt
        card = slide.shapes.add_shape(
            5,  # ROUNDED_RECTANGLE
            Inches(left_inches), Inches(top_inches),
            Inches(width_inches), Inches(height_inches)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = hex_to_pptx_rgb(QdrantColors.BG_ELEVATION_1)
        card.line.color.rgb = hex_to_pptx_rgb(QdrantColors.NEUTRAL_700)
        card.line.width = Pt(1)
        return card

    def save(self, path: str):
        """Save the presentation."""
        self.prs.save(path)


# =============================================================================
# DOCX HELPER
# =============================================================================

class QdrantDocxHelper:
    """Helper for creating Qdrant-branded Word documents."""

    def __init__(self):
        from docx import Document
        from docx.shared import Inches

        self.doc = Document()

        # Set margins
        for section in self.doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        self._configure_styles()

    def _configure_styles(self):
        """Configure document styles to match Qdrant brand."""
        from docx.shared import Pt

        # Normal style
        style = self.doc.styles['Normal']
        font = style.font
        font.name = QdrantTypography.FONT_OFFICE_FALLBACK
        font.size = Pt(QdrantTypography.DOCX_BODY)
        font.color.rgb = hex_to_docx_rgb(QdrantColors.NEUTRAL_800)

        # Heading styles
        heading_sizes = {
            1: QdrantTypography.DOCX_TITLE,
            2: QdrantTypography.DOCX_H1,
            3: QdrantTypography.DOCX_H2,
            4: QdrantTypography.DOCX_H3,
        }
        for level, size in heading_sizes.items():
            hs = self.doc.styles[f'Heading {level}']
            hs.font.name = QdrantTypography.FONT_OFFICE_FALLBACK
            hs.font.size = Pt(size)
            hs.font.bold = True
            hs.font.color.rgb = hex_to_docx_rgb(QdrantColors.NEUTRAL_950)

    def add_heading(self, text: str, level: int = 1):
        """Add a styled heading."""
        return self.doc.add_heading(text, level=level)

    def add_paragraph(self, text: str):
        """Add a styled body paragraph."""
        return self.doc.add_paragraph(text)

    def save(self, path: str):
        """Save the document."""
        self.doc.save(path)


# =============================================================================
# XLSX HELPER
# =============================================================================

class QdrantXlsxHelper:
    """Helper for creating Qdrant-branded spreadsheets."""

    def __init__(self):
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        self.wb = openpyxl.Workbook()
        self.ws = self.wb.active

        # Pre-built styles
        self.header_fill = PatternFill(
            start_color=hex_to_openpyxl(QdrantColors.BG_DEFAULT),
            end_color=hex_to_openpyxl(QdrantColors.BG_DEFAULT),
            fill_type="solid"
        )
        self.header_font = Font(
            name=QdrantTypography.FONT_OFFICE_FALLBACK,
            size=11, bold=True,
            color=hex_to_openpyxl(QdrantColors.NEUTRAL_50)
        )
        self.accent_fill = PatternFill(
            start_color=hex_to_openpyxl(QdrantColors.PRIMARY_500),
            end_color=hex_to_openpyxl(QdrantColors.PRIMARY_500),
            fill_type="solid"
        )
        self.accent_font = Font(
            name=QdrantTypography.FONT_OFFICE_FALLBACK,
            size=11, bold=True, color="FFFFFF"
        )
        self.data_font = Font(
            name=QdrantTypography.FONT_OFFICE_FALLBACK,
            size=10,
            color=hex_to_openpyxl(QdrantColors.NEUTRAL_800)
        )
        self.even_fill = PatternFill(
            start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
        )
        self.odd_fill = PatternFill(
            start_color=hex_to_openpyxl(QdrantColors.BG_LIGHT_SURFACE),
            end_color=hex_to_openpyxl(QdrantColors.BG_LIGHT_SURFACE),
            fill_type="solid"
        )
        self.border = Border(
            left=Side(style="thin", color=hex_to_openpyxl(QdrantColors.NEUTRAL_200)),
            right=Side(style="thin", color=hex_to_openpyxl(QdrantColors.NEUTRAL_200)),
            top=Side(style="thin", color=hex_to_openpyxl(QdrantColors.NEUTRAL_200)),
            bottom=Side(style="thin", color=hex_to_openpyxl(QdrantColors.NEUTRAL_200)),
        )

    def write_header_row(self, headers: list, accent_columns: list = None):
        """Write a branded header row.

        Args:
            headers: List of header strings
            accent_columns: List of 0-based column indices to highlight in Qdrant Red
        """
        from openpyxl.styles import Alignment
        accent_columns = accent_columns or []

        for col, header in enumerate(headers):
            cell = self.ws.cell(row=1, column=col + 1, value=header)
            if col in accent_columns:
                cell.fill = self.accent_fill
                cell.font = self.accent_font
            else:
                cell.fill = self.header_fill
                cell.font = self.header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = self.border

    def write_data_rows(self, data: list, start_row: int = 2):
        """Write data rows with alternating fills.

        Args:
            data: List of lists (rows of values)
            start_row: Row number to start writing (default 2, after header)
        """
        from openpyxl.styles import Alignment

        for row_idx, row_data in enumerate(data):
            fill = self.even_fill if row_idx % 2 == 0 else self.odd_fill
            for col_idx, value in enumerate(row_data):
                cell = self.ws.cell(
                    row=start_row + row_idx,
                    column=col_idx + 1,
                    value=value
                )
                cell.font = self.data_font
                cell.fill = fill
                cell.border = self.border
                cell.alignment = Alignment(
                    horizontal='center' if col_idx > 0 else 'left'
                )

    def auto_width(self):
        """Auto-adjust column widths based on content."""
        for col in self.ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            self.ws.column_dimensions[col_letter].width = min(max_length + 4, 40)

    def save(self, path: str):
        """Save the workbook."""
        self.auto_width()
        self.wb.save(path)
